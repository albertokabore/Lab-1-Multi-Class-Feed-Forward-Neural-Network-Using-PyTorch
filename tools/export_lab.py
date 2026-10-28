"""Build clean Word and PDF submissions from the executed notebook."""

import base64
import html
import io
import re
import textwrap
from pathlib import Path

import mistune
import nbformat
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image, PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer,
    Table, TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
STEM = ROOT / "Lab1_FFNN_BoW_AG_News"
MARKDOWN = mistune.create_markdown(renderer="ast", plugins=["table"])
FONT_DIR = Path("C:/Windows/Fonts")


def plain(nodes):
    return "".join(
        plain(node["children"]) if "children" in node else
        " " if node["type"] == "softbreak" else
        "\n" if node["type"] in {"linebreak", "inline_html"} else
        node.get("raw", "")
        for node in nodes
    )


def rich(nodes):
    result = []
    for node in nodes:
        kind = node["type"]
        inner = rich(node.get("children", []))
        if kind == "strong":
            result.append(f"<b>{inner}</b>")
        elif kind == "emphasis":
            result.append(f"<i>{inner}</i>")
        elif kind == "codespan":
            result.append(f'<font name="LabMono">{html.escape(node["raw"])}</font>')
        elif kind == "link":
            url = html.escape(node["attrs"]["url"], quote=True)
            result.append(f'<link href="{url}" color="#245b82">{inner}</link>')
        elif kind in {"inline_html", "linebreak"}:
            result.append("<br/>")
        elif kind == "softbreak":
            result.append(" ")
        elif "children" in node:
            result.append(inner)
        else:
            result.append(html.escape(node.get("raw", "")))
    return "".join(result)


def word_inline(paragraph, nodes, bold=False, italic=False):
    for node in nodes:
        kind = node["type"]
        if "children" in node:
            word_inline(paragraph, node["children"], bold or kind == "strong",
                        italic or kind == "emphasis")
            if kind == "link":
                paragraph.add_run(f' ({node["attrs"]["url"]})')
            continue
        value = " " if kind == "softbreak" else (
            "\n" if kind in {"linebreak", "inline_html"} else node.get("raw", ""))
        run = paragraph.add_run(value)
        run.bold, run.italic = bold, italic
        if kind == "codespan":
            run.font.name = "Consolas"
            run.font.size = Pt(9)


class Exporter:
    def __init__(self):
        for name, file in [("LabText", "calibri.ttf"), ("LabBold", "calibrib.ttf"),
                           ("LabItalic", "calibrii.ttf"), ("LabMono", "consola.ttf")]:
            pdfmetrics.registerFont(TTFont(name, str(FONT_DIR / file)))
        pdfmetrics.registerFontFamily("LabText", normal="LabText", bold="LabBold",
                                      italic="LabItalic", boldItalic="LabBold")
        self.doc = Document()
        section = self.doc.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.65)
        section.left_margin = section.right_margin = Inches(0.65)
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = "Calibri", Pt(10)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.08
        for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
            self.doc.styles[name].font.name = "Calibri"
            self.doc.styles[name].font.color.rgb = RGBColor.from_string("17365D")
        self.doc.core_properties.title = "Lab 1: Multi-Class Feed-Forward Neural Network Using PyTorch"
        self.doc.core_properties.author = "Albert Kabore"
        self.doc.core_properties.keywords = "AG News, PyTorch, bag of words, text classification"
        footer = section.footer.paragraphs[0]
        footer.alignment = 2
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        self.styles = getSampleStyleSheet()
        for name in ["Normal", "BodyText"]:
            self.styles[name].fontName = "LabText"
            self.styles[name].fontSize = 10
            self.styles[name].leading = 13
            self.styles[name].spaceAfter = 7
            self.styles[name].alignment = TA_LEFT
            self.styles[name].splitLongWords = True
        for name in ["Title", "Heading1", "Heading2", "Heading3"]:
            self.styles[name].fontName = "LabBold"
            self.styles[name].textColor = colors.HexColor("#17365D")
        self.styles.add(ParagraphStyle("LabCode", fontName="LabMono", fontSize=8,
                                        leading=10.2, spaceBefore=3, spaceAfter=7))
        self.styles.add(ParagraphStyle("LabCell", parent=self.styles["BodyText"],
                                        fontSize=8.3, leading=10.5, spaceAfter=0))
        self.story = []
        self.width = 7.2 * 72
        self.figure_number = 0

    def code(self, text, output=False):
        text = text.replace("\r", "").strip("\n")
        if not text:
            return
        # Wrap only the rendered exports; executable notebook source stays intact.
        lines = []
        for line in text.splitlines():
            lines.extend(textwrap.wrap(line.expandtabs(4), width=112,
                                       replace_whitespace=False, drop_whitespace=False,
                                       subsequent_indent="    ") or [""])
        for line in lines:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(line or " ")
            run.font.name, run.font.size = "Consolas", Pt(8)
            if output:
                run.font.color.rgb = RGBColor.from_string("354355")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        self.story.append(Preformatted("\n".join(lines), self.styles["LabCode"]))

    def table(self, node):
        rows = []
        for child in node["children"]:
            row_nodes = [child] if child["type"] == "table_head" else child["children"]
            for row in row_nodes:
                rows.append([plain(cell["children"]) for cell in row["children"]])
        table = self.doc.add_table(rows=0, cols=len(rows[0]))
        table.style = "Light Shading Accent 1"
        for values in rows:
            for cell, value in zip(table.add_row().cells, values):
                cell.text = value
        pdfrows = [[Paragraph(html.escape(value), self.styles["LabCell"])
                    for value in row] for row in rows]
        pdf_table = Table(pdfrows, colWidths=[self.width / len(rows[0])] * len(rows[0]),
                          repeatRows=1, hAlign="LEFT")
        pdf_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e7edf4")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cad2dc")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        self.story.extend([pdf_table, Spacer(1, 8)])

    def blocks(self, nodes):
        for node in nodes:
            kind = node["type"]
            if kind == "heading":
                level = node["attrs"]["level"]
                text = plain(node["children"])
                if level == 2 and not text.startswith("Part A:"):
                    self.doc.add_page_break()
                    self.story.append(PageBreak())
                self.doc.add_heading(text, level=0 if level == 1 else min(level - 1, 3))
                style = "Title" if level == 1 else f"Heading{min(level - 1, 3)}"
                self.story.append(Paragraph(rich(node["children"]), self.styles[style]))
            elif kind in {"paragraph", "block_text"}:
                word_inline(self.doc.add_paragraph(), node["children"])
                self.story.append(Paragraph(rich(node["children"]), self.styles["BodyText"]))
            elif kind == "list":
                ordered = node.get("attrs", {}).get("ordered", False)
                for i, item in enumerate(node["children"], 1):
                    content = item["children"][0].get("children", [])
                    word_inline(self.doc.add_paragraph(style="List Number" if ordered else "List Bullet"), content)
                    prefix = f"{i}. " if ordered else "• "
                    self.story.append(Paragraph(prefix + rich(content), self.styles["BodyText"]))
            elif kind == "table":
                self.table(node)
            elif kind == "block_code":
                self.code(node["raw"])

    def figure(self, encoded):
        self.figure_number += 1
        data = base64.b64decode(encoded)
        self.doc.add_picture(io.BytesIO(data), width=Inches(6.5))
        im = Image(io.BytesIO(data))
        ratio = min(self.width / im.imageWidth, 330 / im.imageHeight, 1)
        im.drawWidth, im.drawHeight = im.imageWidth * ratio, im.imageHeight * ratio
        im.hAlign = "LEFT"
        self.story.extend([im, Spacer(1, 8)])

    def export(self, notebook):
        errors = [o for c in notebook.cells for o in c.get("outputs", [])
                  if o.output_type == "error"]
        if errors or any(c.execution_count is None for c in notebook.cells if c.cell_type == "code"):
            raise ValueError("Execute the complete notebook successfully before exporting.")
        for cell in notebook.cells:
            if cell.cell_type == "markdown":
                self.blocks(MARKDOWN(cell.source))
                continue
            if cell.cell_type != "code":
                continue
            self.code(cell.source)
            for output in cell.get("outputs", []):
                data = output.get("data", {})
                if "image/png" in data:
                    self.figure(data["image/png"])
                elif "application/vnd.jupyter.widget-view+json" in data:
                    continue
                elif output.output_type == "stream":
                    # Progress-bar redraws are terminal controls, not report results.
                    value = output.get("text", "")
                    if output.get("name") != "stderr":
                        self.code(value, output=True)
                elif "text/plain" in data:
                    value = data["text/plain"]
                    if not re.match(r"<(?:Figure|Axes)", value):
                        self.code(value, output=True)
        self.doc.save(STEM.with_suffix(".docx"))

        def page_number(canvas, document):
            canvas.setFont("LabText", 8)
            canvas.setFillColor(colors.HexColor("#64748b"))
            canvas.drawRightString(7.85 * 72, 0.35 * 72, str(document.page))

        pdf = SimpleDocTemplate(str(STEM.with_suffix(".pdf")), pagesize=(612, 792),
                                rightMargin=46.8, leftMargin=46.8,
                                topMargin=46.8, bottomMargin=46.8,
                                title=self.doc.core_properties.title, author="Albert Kabore")
        pdf.build(self.story, onFirstPage=page_number, onLaterPages=page_number)
        print(f"Exported Word and PDF with {self.figure_number} figures.")


if __name__ == "__main__":
    Exporter().export(nbformat.read(STEM.with_suffix(".ipynb"), as_version=4))
